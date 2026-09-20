import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image
from io import BytesIO
from pathlib import Path
from datetime import datetime
import tempfile

st.set_page_config(page_title="PSO Outlet Improvement Report", page_icon="📊", layout="centered")

ASSETS = Path(__file__).parent / "assets"
BLUE = RGBColor(0, 82, 165)
GREEN = RGBColor(0, 154, 68)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(245, 248, 250)
WHITE = RGBColor(255,255,255)

if "step" not in st.session_state:
    st.session_state.step = 0
if "outlet" not in st.session_state:
    st.session_state.outlet = {"name":"", "code":"", "presented_by":"", "designation":""}
if "items" not in st.session_state:
    st.session_state.items = []

st.markdown("""
<style>
.block-container {max-width: 1050px; padding-top: 1.2rem;}
.pso-title {font-size: 2rem; font-weight: 800; color:#0052A5; margin-bottom:0.1rem;}
.pso-sub {color:#4b5563; margin-bottom:1.2rem;}
.step-card {padding:1rem 1.2rem; border:1px solid #dfe5eb; border-radius:14px; background:#fff;}
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="pso-title">PSO Outlet Improvement Report</div>', unsafe_allow_html=True)
st.markdown('<div class="pso-sub">Create a professional Before & After PowerPoint presentation from your outlet improvement photographs.</div>', unsafe_allow_html=True)


def safe_filename(s):
    return "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in s).strip().replace(" ", "_") or "Outlet"

def add_text(slide, text, x,y,w,h, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(x,y,w,h)
    tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def add_header(slide, outlet_name, title=None, slide_no=None):
    # White header with PSO logo on left and outlet on right.
    logo_path=ASSETS/"pso_logo.png"
    slide.shapes.add_picture(str(logo_path), Inches(0.35), Inches(0.18), height=Inches(0.52))
    add_text(slide, outlet_name, Inches(7.0), Inches(0.22), Inches(6.0), Inches(0.4), size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    # green/blue rule
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(13.333), Inches(0.045))
    line.fill.solid(); line.fill.fore_color.rgb=GREEN; line.line.fill.background()
    line2=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(8.9), Inches(0.045))
    line2.fill.solid(); line2.fill.fore_color.rgb=BLUE; line2.line.fill.background()
    if slide_no is not None:
        add_text(slide, str(slide_no), Inches(12.55), Inches(7.08), Inches(0.4), Inches(0.25), size=9, color=RGBColor(100,100,100), align=PP_ALIGN.RIGHT)

def add_picture_cover(slide, data, x,y,w,h):
    # Crop image to fill target rectangle while preserving aspect ratio.
    img=Image.open(BytesIO(data)).convert("RGB")
    tw,th=int(w),int(h)
    iw,ih=img.size
    target=tw/th
    ratio=iw/ih
    if ratio>target:
        nw=int(ih*target); left=(iw-nw)//2; img=img.crop((left,0,left+nw,ih))
    else:
        nh=int(iw/target); top=(ih-nh)//2; img=img.crop((0,top,iw,top+nh))
    out=BytesIO(); img.save(out,format="JPEG",quality=94); out.seek(0)
    slide.shapes.add_picture(out, x,y,w,h)

def add_desc(slide, text, x,y,w,h):
    box=slide.shapes.add_textbox(x,y,w,h)
    tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(0.06); tf.margin_right=Inches(0.06); tf.margin_top=Inches(0.04)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=text; r.font.name="Aptos"; r.font.size=Pt(12); r.font.color.rgb=DARK
    return box

def build_ppt(outlet, items):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    # Cover
    slide=prs.slides.add_slide(blank)
    # Use supplied template cover visual as full background.
    slide.shapes.add_picture(str(ASSETS/"cover.jpg"), 0,0, width=prs.slide_width, height=prs.slide_height)
    # Required header branding: PSO logo on every slide and outlet name at top-right.
    slide.shapes.add_picture(str(ASSETS/"pso_logo.png"), Inches(0.35), Inches(0.18), height=Inches(0.62))
    add_text(slide, outlet["name"], Inches(7.0), Inches(0.22), Inches(5.95), Inches(0.4), size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    # translucent-ish white panel
    panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.45), Inches(6.1), Inches(3.2))
    panel.fill.solid(); panel.fill.fore_color.rgb=WHITE; panel.fill.transparency=8; panel.line.fill.background()
    add_text(slide, outlet["name"], Inches(1.0), Inches(1.95), Inches(5.4), Inches(0.7), size=28, bold=True, color=DARK)
    add_text(slide, outlet["code"], Inches(1.0), Inches(2.72), Inches(5.4), Inches(0.55), size=22, bold=True, color=DARK)
    add_text(slide, "OUTLET IMPROVEMENT REPORT", Inches(1.0), Inches(3.38), Inches(5.4), Inches(0.6), size=18, bold=True, color=BLUE)
    if outlet.get("presented_by") or outlet.get("designation"):
        txt=""
        if outlet.get("presented_by"): txt += "Presented by: " + outlet["presented_by"] + "\n"
        if outlet.get("designation"): txt += "Designation: " + outlet["designation"]
        add_text(slide, txt, Inches(1.0), Inches(4.05), Inches(5.4), Inches(0.7), size=11, color=DARK)
    # Improvement slides
    for idx,item in enumerate(items,1):
        slide=prs.slides.add_slide(blank)
        # background
        bg=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb=WHITE; bg.line.fill.background()
        add_header(slide,outlet["name"],item["heading"],idx)
        add_text(slide, f"{idx}. {item['heading']}", Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.55), size=23, bold=True, color=BLUE)
        # Cards
        left_x=Inches(0.55); right_x=Inches(6.85); top=Inches(1.72); img_w=Inches(5.95); img_h=Inches(3.05)
        for x,label,data,desc in [(left_x,"BEFORE",item["before"],item["before_desc"]),(right_x,"AFTER",item["after"],item["after_desc"])]:
            card=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,top,img_w,Inches(4.95))
            card.fill.solid(); card.fill.fore_color.rgb=LIGHT; card.line.color.rgb=RGBColor(220,226,232)
            add_text(slide,label,x+Inches(0.15),top+Inches(0.12),Inches(1.2),Inches(0.35),size=12,bold=True,color=GREEN)
            add_picture_cover(slide,data,x+Inches(0.15),top+Inches(0.52),img_w-Inches(0.3),img_h)
            add_desc(slide,desc,x+Inches(0.15),top+Inches(3.72),img_w-Inches(0.3),Inches(1.0))
        add_text(slide,"Pakistan State Oil Company Limited",Inches(0.55),Inches(7.02),Inches(6.0),Inches(0.25),size=8,bold=True,color=RGBColor(90,90,90))
    out=BytesIO(); prs.save(out); out.seek(0); return out.getvalue()

if st.session_state.step == 0:
    st.subheader("1. Outlet Information")
    with st.form("outlet_form"):
        name = st.text_input("Outlet Name *", value=st.session_state.outlet["name"], placeholder="e.g. Clifton Service Station")
        code = st.text_input("Outlet Code *", value=st.session_state.outlet["code"], placeholder="e.g. 1234")
        st.markdown("**Optional presentation details**")
        presented_by = st.text_input("Presented By", value=st.session_state.outlet["presented_by"])
        designation = st.text_input("Designation", value=st.session_state.outlet["designation"])
        start = st.form_submit_button("Start Improvement Report →", type="primary", use_container_width=True)
    if start:
        if not name.strip() or not code.strip():
            st.error("Please enter both Outlet Name and Outlet Code.")
        else:
            st.session_state.outlet = {"name":name.strip(), "code":code.strip(), "presented_by":presented_by.strip(), "designation":designation.strip()}
            st.session_state.items=[]
            st.session_state.step=1
            st.rerun()

elif st.session_state.step == 1:
    n = len(st.session_state.items)+1
    st.subheader(f"2. Improvement {n}")
    st.caption(f"Outlet: {st.session_state.outlet['name']}  |  Code: {st.session_state.outlet['code']}")

    with st.form(f"improvement_{n}"):
        heading = st.text_input("Improvement Heading *", placeholder=f"e.g. Forecourt Painting / Canopy Enhancement")
        c1,c2 = st.columns(2)
        with c1:
            st.markdown("### Before")
            before = st.file_uploader("Upload Before Picture *", type=["jpg","jpeg","png","webp"], key=f"before_{n}")
            before_desc = st.text_area("Before Description *", height=120, placeholder="Describe the condition before improvement...")
        with c2:
            st.markdown("### After")
            after = st.file_uploader("Upload After Picture *", type=["jpg","jpeg","png","webp"], key=f"after_{n}")
            after_desc = st.text_area("After Description *", height=120, placeholder="Describe the improvement after completion...")

        st.markdown("---")
        b1,b2 = st.columns(2)
        with b1:
            cont = st.form_submit_button("➕ Save & Continue", use_container_width=True)
        with b2:
            end = st.form_submit_button("✅ Save & Generate PPT", type="primary", use_container_width=True)

    action = "continue" if cont else ("end" if end else None)
    if action:
        errors=[]
        if not heading.strip(): errors.append("Improvement heading")
        if before is None: errors.append("Before picture")
        if not before_desc.strip(): errors.append("Before description")
        if after is None: errors.append("After picture")
        if not after_desc.strip(): errors.append("After description")
        if errors:
            st.error("Please complete: " + ", ".join(errors) + ".")
        else:
            st.session_state.items.append({
                "heading":heading.strip(),
                "before":before.getvalue(),
                "before_name":before.name,
                "before_desc":before_desc.strip(),
                "after":after.getvalue(),
                "after_name":after.name,
                "after_desc":after_desc.strip(),
            })
            if action == "continue":
                st.rerun()
            else:
                st.session_state.step=2
                st.rerun()

    if st.session_state.items:
        st.markdown("### Improvements already added")
        for i,item in enumerate(st.session_state.items,1):
            st.write(f"**{i}. {item['heading']}**")
        if st.button("↩ Start Over", use_container_width=True):
            st.session_state.clear()
            st.rerun()

elif st.session_state.step == 2:
    st.subheader("3. Presentation Ready")
    st.success(f"{len(st.session_state.items)} improvement(s) captured for {st.session_state.outlet['name']}.")
    st.write("Click the button below to generate the PowerPoint presentation.")
    if st.button("🎨 Generate Professional PowerPoint", type="primary", use_container_width=True):
        pptx_bytes = build_ppt(st.session_state.outlet, st.session_state.items)
        filename = f"{safe_filename(st.session_state.outlet['name'])}_Improvement_Report.pptx"
        st.download_button("⬇️ Download PowerPoint", data=pptx_bytes, file_name=filename, mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", use_container_width=True)
        st.info("Your presentation has been generated. You can download it above.")
    if st.button("＋ Create Another Report", use_container_width=True):
        st.session_state.clear()
        st.rerun()

