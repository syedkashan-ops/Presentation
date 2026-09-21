import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageOps
from io import BytesIO
from pathlib import Path

st.set_page_config(page_title="PSO Outlet Improvement Report", page_icon="📊", layout="centered")

ASSETS = Path(__file__).parent / "assets"
BLUE = RGBColor(0, 82, 165)
GREEN = RGBColor(0, 154, 68)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(245, 248, 250)
WHITE = RGBColor(255, 255, 255)

MAX_IMAGE_SIDE = 1600
JPEG_QUALITY = 82

# ---------------- Session state ----------------
def init_state():
    defaults = {
        "step": 0,
        "outlet": {"name": "", "code": "", "presented_by": "", "designation": ""},
        "improvements": [],
        "current_uploads": {},
        "photo_stage": "before",
        "draft_heading": "",
        "draft_before_desc": "",
        "draft_after_desc": "",
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v

init_state()

st.markdown("""
<style>
.block-container {max-width: 1050px; padding-top: 1.2rem;}
.pso-title {font-size: 2rem; font-weight: 800; color:#0052A5; margin-bottom:0.1rem;}
.pso-sub {color:#4b5563; margin-bottom:1.2rem;}
.small-note {font-size:0.88rem; color:#5b6470;}
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="pso-title">PSO Outlet Improvement Report</div>', unsafe_allow_html=True)
st.markdown('<div class="pso-sub">Create a professional Before & After PowerPoint presentation from your outlet improvement photographs.</div>', unsafe_allow_html=True)


def safe_filename(s):
    return "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in s).strip().replace(" ", "_") or "Outlet"


def compress_image(uploaded_bytes):
    """Validate, orient, resize and compress a phone image before storing it."""
    img = Image.open(BytesIO(uploaded_bytes))
    img = ImageOps.exif_transpose(img)
    img = img.convert("RGB")
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE), Image.Resampling.LANCZOS)
    out = BytesIO()
    img.save(out, format="JPEG", quality=JPEG_QUALITY, optimize=True, progressive=True)
    data = out.getvalue()
    return data, img.size


def save_uploaded_photo(uploaded_file, key):
    raw = uploaded_file.getvalue()
    if not raw:
        raise ValueError("The selected file is empty.")
    compressed, size = compress_image(raw)
    st.session_state.current_uploads[key] = {
        "data": compressed,
        "name": uploaded_file.name,
        "size": size,
        "original_bytes": len(raw),
        "compressed_bytes": len(compressed),
    }


def photo_size_text(record):
    if not record:
        return ""
    old_mb = record.get("original_bytes", 0) / 1024 / 1024
    new_mb = record.get("compressed_bytes", 0) / 1024 / 1024
    w, h = record.get("size", (0, 0))
    return f"Optimized to {w}×{h}px • {old_mb:.1f} MB → {new_mb:.1f} MB"


def add_text(slide, text, x, y, w, h, size=18, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_header(slide, outlet_name, slide_no=None):
    slide.shapes.add_picture(str(ASSETS / "pso_logo.png"), Inches(0.35), Inches(0.18), height=Inches(0.52))
    add_text(slide, outlet_name, Inches(7.0), Inches(0.22), Inches(6.0), Inches(0.4),
             size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(13.333), Inches(0.045))
    line.fill.solid(); line.fill.fore_color.rgb = GREEN; line.line.fill.background()
    line2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(8.9), Inches(0.045))
    line2.fill.solid(); line2.fill.fore_color.rgb = BLUE; line2.line.fill.background()
    if slide_no is not None:
        add_text(slide, str(slide_no), Inches(12.55), Inches(7.08), Inches(0.4), Inches(0.25),
                 size=9, color=RGBColor(100, 100, 100), align=PP_ALIGN.RIGHT)


def add_picture_cover(slide, data, x, y, w, h):
    img = Image.open(BytesIO(data)).convert("RGB")
    tw, th = int(w), int(h)
    iw, ih = img.size
    target = tw / th
    ratio = iw / ih
    if ratio > target:
        nw = int(ih * target)
        left = max((iw - nw) // 2, 0)
        img = img.crop((left, 0, left + nw, ih))
    else:
        nh = int(iw / target)
        top = max((ih - nh) // 2, 0)
        img = img.crop((0, top, iw, top + nh))
    out = BytesIO()
    img.save(out, format="JPEG", quality=90, optimize=True)
    out.seek(0)
    slide.shapes.add_picture(out, x, y, w, h)


def add_desc(slide, text, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06); tf.margin_top = Inches(0.04)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.name = "Aptos"; r.font.size = Pt(12); r.font.color.rgb = DARK
    return box


def build_ppt(outlet, improvements):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(ASSETS / "cover.jpg"), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.shapes.add_picture(str(ASSETS / "pso_logo.png"), Inches(0.35), Inches(0.18), height=Inches(0.62))
    add_text(slide, outlet["name"], Inches(7.0), Inches(0.22), Inches(5.95), Inches(0.4),
             size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.45), Inches(6.1), Inches(3.2))
    panel.fill.solid(); panel.fill.fore_color.rgb = WHITE; panel.fill.transparency = 8; panel.line.fill.background()
    add_text(slide, outlet["name"], Inches(1.0), Inches(1.95), Inches(5.4), Inches(0.7), size=28, bold=True)
    add_text(slide, outlet["code"], Inches(1.0), Inches(2.72), Inches(5.4), Inches(0.55), size=22, bold=True)
    add_text(slide, "OUTLET IMPROVEMENT REPORT", Inches(1.0), Inches(3.38), Inches(5.4), Inches(0.6), size=18, bold=True, color=BLUE)
    if outlet.get("presented_by") or outlet.get("designation"):
        txt = ""
        if outlet.get("presented_by"): txt += "Presented by: " + outlet["presented_by"] + "\n"
        if outlet.get("designation"): txt += "Designation: " + outlet["designation"]
        add_text(slide, txt, Inches(1.0), Inches(4.05), Inches(5.4), Inches(0.7), size=11)

    for idx, item in enumerate(improvements, 1):
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background()
        add_header(slide, outlet["name"], idx)
        add_text(slide, f"{idx}. {item['heading']}", Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.55),
                 size=23, bold=True, color=BLUE)
        left_x = Inches(0.55); right_x = Inches(6.85); top = Inches(1.72)
        img_w = Inches(5.95); img_h = Inches(3.05)
        for x, label, data, desc in [
            (left_x, "BEFORE", item["before"], item["before_desc"]),
            (right_x, "AFTER", item["after"], item["after_desc"]),
        ]:
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, img_w, Inches(4.95))
            card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = RGBColor(220, 226, 232)
            add_text(slide, label, x + Inches(0.15), top + Inches(0.12), Inches(1.2), Inches(0.35), size=12, bold=True, color=GREEN)
            add_picture_cover(slide, data, x + Inches(0.15), top + Inches(0.52), img_w - Inches(0.3), img_h)
            add_desc(slide, desc, x + Inches(0.15), top + Inches(3.72), img_w - Inches(0.3), Inches(1.0))
        add_text(slide, "Pakistan State Oil Company Limited", Inches(0.55), Inches(7.02), Inches(6.0), Inches(0.25),
                 size=8, bold=True, color=RGBColor(90, 90, 90))

    out = BytesIO(); prs.save(out); out.seek(0)
    return out.getvalue()


def reset_current_draft():
    st.session_state.current_uploads = {}
    st.session_state.photo_stage = "before"
    st.session_state.draft_heading = ""
    st.session_state.draft_before_desc = ""
    st.session_state.draft_after_desc = ""


def current_keys(n):
    return f"before_{n}", f"after_{n}"


# ---------------- Outlet ----------------
if st.session_state.step == 0:
    st.subheader("1. Outlet Information")
    with st.form("outlet_form"):
        name = st.text_input("Outlet Name *", value=st.session_state.outlet["name"], placeholder="e.g. Clifton Service Station")
        code = st.text_input("Outlet Code *", value=st.session_state.outlet["code"], placeholder="e.g. 1234")
        st.markdown("**Optional presentation details**")
        presented_by = st.text_input("Presented By", value=st.session_state.outlet["presented_by"])
        designation = st.text_input("Designation", value=st.session_state.outlet["designation"])
        start = st.form_submit_button("Start Improvement Report →", type="primary", use_container_width=True)
    if start:
        if not name.strip() or not code.strip():
            st.error("Please enter both Outlet Name and Outlet Code.")
        else:
            st.session_state.outlet = {
                "name": name.strip(), "code": code.strip(),
                "presented_by": presented_by.strip(), "designation": designation.strip()
            }
            st.session_state.improvements = []
            reset_current_draft()
            st.session_state.step = 1
            st.rerun()

# ---------------- Improvement capture ----------------
elif st.session_state.step == 1:
    n = len(st.session_state.improvements) + 1
    st.subheader(f"2. Improvement {n}")
    st.caption(f"Outlet: {st.session_state.outlet['name']}  |  Code: {st.session_state.outlet['code']}")

    heading = st.text_input(
        "Improvement Heading *",
        key="draft_heading",
        placeholder="e.g. Forecourt Painting / Canopy Enhancement"
    )

    before_key, after_key = current_keys(n)

    # BEFORE
    if st.session_state.photo_stage == "before":
        st.markdown("### Before Picture")
        st.info("Select one photo from your phone Gallery. The app automatically compresses it after upload to make the PowerPoint smaller.")
        before = st.file_uploader(
            "Upload Before Picture *",
            type=["jpg", "jpeg", "png", "webp"],
            key=before_key,
            help="For best mobile performance, JPG photos are recommended. Large phone photos are automatically compressed."
        )
        if before is not None:
            try:
                save_uploaded_photo(before, before_key)
                rec = st.session_state.current_uploads[before_key]
                st.image(rec["data"], caption="Before picture", use_container_width=True)
                st.success("Before picture uploaded and compressed.")
                st.caption(photo_size_text(rec))
                if st.button("Continue to Before Description →", type="primary", use_container_width=True):
                    st.session_state.photo_stage = "before_description"
                    st.rerun()
            except Exception as e:
                st.error(f"The selected file could not be processed. Please select another JPG/PNG photo. ({type(e).__name__})")
        elif before_key in st.session_state.current_uploads:
            rec = st.session_state.current_uploads[before_key]
            st.image(rec["data"], caption="Before picture", use_container_width=True)
            st.success("Before picture is saved.")
            if st.button("Continue to Before Description →", type="primary", use_container_width=True):
                st.session_state.photo_stage = "before_description"
                st.rerun()

    elif st.session_state.photo_stage == "before_description":
        st.markdown("### Before Description")
        rec = st.session_state.current_uploads.get(before_key)
        if not rec:
            st.session_state.photo_stage = "before"
            st.rerun()
        st.image(rec["data"], caption="Before picture", use_container_width=True)
        st.session_state.draft_before_desc = st.text_area(
            "Before Description *",
            value=st.session_state.draft_before_desc,
            key="draft_before_desc_widget",
            height=130,
            placeholder="Describe the condition before improvement..."
        )
        c1, c2 = st.columns(2)
        with c1:
            if st.button("← Replace Before Photo", use_container_width=True):
                st.session_state.current_uploads.pop(before_key, None)
                st.session_state.photo_stage = "before"
                st.session_state.draft_before_desc = ""
                st.rerun()
        with c2:
            if st.button("Continue to After Photo →", type="primary", use_container_width=True):
                desc = st.session_state.get("draft_before_desc_widget", st.session_state.draft_before_desc).strip()
                st.session_state.draft_before_desc = desc
                if not heading.strip():
                    st.error("Please enter the Improvement Heading.")
                elif not desc:
                    st.error("Please enter the Before Description.")
                else:
                    st.session_state.photo_stage = "after"
                    st.rerun()

    # AFTER
    elif st.session_state.photo_stage == "after":
        st.markdown("### After Picture")
        st.info("Select the After photo from your Gallery. The Before photo and description are already saved.")
        after = st.file_uploader(
            "Upload After Picture *",
            type=["jpg", "jpeg", "png", "webp"],
            key=after_key,
            help="Large phone photos are automatically compressed."
        )
        if after is not None:
            try:
                save_uploaded_photo(after, after_key)
                rec = st.session_state.current_uploads[after_key]
                st.image(rec["data"], caption="After picture", use_container_width=True)
                st.success("After picture uploaded and compressed.")
                st.caption(photo_size_text(rec))
                if st.button("Continue to After Description →", type="primary", use_container_width=True):
                    st.session_state.photo_stage = "after_description"
                    st.rerun()
            except Exception as e:
                st.error(f"The selected file could not be processed. Please select another JPG/PNG photo. ({type(e).__name__})")
        elif after_key in st.session_state.current_uploads:
            rec = st.session_state.current_uploads[after_key]
            st.image(rec["data"], caption="After picture", use_container_width=True)
            st.success("After picture is saved.")
            if st.button("Continue to After Description →", type="primary", use_container_width=True):
                st.session_state.photo_stage = "after_description"
                st.rerun()

    # AFTER DESCRIPTION / SAVE
    elif st.session_state.photo_stage == "after_description":
        st.markdown("### After Description")
        rec = st.session_state.current_uploads.get(after_key)
        if not rec:
            st.session_state.photo_stage = "after"
            st.rerun()
        st.image(rec["data"], caption="After picture", use_container_width=True)
        st.session_state.draft_after_desc = st.text_area(
            "After Description *",
            value=st.session_state.draft_after_desc,
            key="draft_after_desc_widget",
            height=130,
            placeholder="Describe the improvement after completion..."
        )

        st.markdown("---")
        st.caption("Before the presentation is generated, the app checks the heading, both descriptions and both photos again.")
        c1, c2 = st.columns(2)
        with c1:
            if st.button("← Replace After Photo", use_container_width=True):
                st.session_state.current_uploads.pop(after_key, None)
                st.session_state.photo_stage = "after"
                st.session_state.draft_after_desc = ""
                st.rerun()

        def validate_current():
            # Always read the actual current widget values; do not rely on a stale local variable.
            h = st.session_state.get("draft_heading", "").strip()
            bd = st.session_state.get("draft_before_desc_widget", st.session_state.get("draft_before_desc", "")).strip()
            ad = st.session_state.get("draft_after_desc_widget", st.session_state.get("draft_after_desc", "")).strip()
            b = st.session_state.current_uploads.get(before_key)
            a = st.session_state.current_uploads.get(after_key)
            return h, bd, ad, b, a

        with c2:
            if st.button("➕ Save & Continue", use_container_width=True):
                h, bd, ad, b, a = validate_current()
                if not h:
                    st.error("Please enter the Improvement Heading.")
                elif not bd:
                    st.error("Please enter the Before Description.")
                elif not b:
                    st.error("Please upload the Before Picture.")
                elif not a:
                    st.error("Please upload the After Picture.")
                elif not ad:
                    st.error("Please enter the After Description.")
                else:
                    st.session_state.improvements.append({
                        "heading": h,
                        "before": b["data"], "before_name": b["name"], "before_desc": bd,
                        "after": a["data"], "after_name": a["name"], "after_desc": ad,
                    })
                    reset_current_draft()
                    st.rerun()

        if st.button("✅ Save & Generate PPT", type="primary", use_container_width=True):
            h, bd, ad, b, a = validate_current()
            if not h:
                st.error("Please enter the Improvement Heading.")
            elif not bd:
                st.error("Please enter the Before Description.")
            elif not b:
                st.error("Please upload the Before Picture.")
            elif not a:
                st.error("Please upload the After Picture.")
            elif not ad:
                st.error("Please enter the After Description.")
            else:
                st.session_state.improvements.append({
                    "heading": h,
                    "before": b["data"], "before_name": b["name"], "before_desc": bd,
                    "after": a["data"], "after_name": a["name"], "after_desc": ad,
                })
                reset_current_draft()
                st.session_state.step = 2
                st.rerun()

    if st.session_state.improvements:
        st.markdown("### Improvements already added")
        for i, item in enumerate(st.session_state.improvements, 1):
            st.write(f"**{i}. {item['heading']}**")
        if st.button("↩ Start Over", use_container_width=True):
            st.session_state.clear()
            st.rerun()

# ---------------- Generate ----------------
elif st.session_state.step == 2:
    st.subheader("3. Presentation Ready")
    st.success(f"{len(st.session_state.improvements)} improvement(s) captured for {st.session_state.outlet['name']}.")
    st.write("Click the button below to generate the PowerPoint presentation.")
    if st.button("🎨 Generate Professional PowerPoint", type="primary", use_container_width=True):
        try:
            if not st.session_state.improvements:
                st.error("No improvement has been saved yet.")
            else:
                pptx_bytes = build_ppt(st.session_state.outlet, st.session_state.improvements)
                filename = f"{safe_filename(st.session_state.outlet['name'])}_Improvement_Report.pptx"
                st.download_button(
                    "⬇️ Download PowerPoint", data=pptx_bytes, file_name=filename,
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True
                )
                st.info("Your presentation has been generated. You can download it above.")
        except Exception as e:
            st.error(f"PowerPoint generation failed: {type(e).__name__}: {e}")

    if st.button("＋ Create Another Report", use_container_width=True):
        st.session_state.clear()
        st.rerun()
