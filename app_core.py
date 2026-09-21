from io import BytesIO
from pathlib import Path
import re

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BLUE = RGBColor(0, 82, 165)
GREEN = RGBColor(0, 154, 68)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(245, 248, 250)
WHITE = RGBColor(255, 255, 255)

MAX_IMAGE_SIDE = 1600
TARGET_BYTES = int(1.5 * 1024 * 1024)


def safe_filename(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(r"[^A-Za-z0-9 _-]+", "_", text)
    text = re.sub(r"\s+", "_", text)
    return text or "Outlet"


def _encode_jpeg(img: Image.Image, quality: int) -> bytes:
    out = BytesIO()
    img.save(out, format="JPEG", quality=quality, optimize=True, progressive=True)
    return out.getvalue()


def compress_image(raw: bytes):
    """Validate, orient, resize and compress an image for reliable PPT use."""
    if not raw:
        raise ValueError("Selected image is empty.")

    with Image.open(BytesIO(raw)) as opened:
        img = ImageOps.exif_transpose(opened)
        img = img.convert("RGB")

    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE), Image.Resampling.LANCZOS)

    data = b""
    for quality in (84, 78, 72, 66, 60, 54):
        data = _encode_jpeg(img, quality)
        if len(data) <= TARGET_BYTES:
            break

    # Very detailed photos can still exceed the target. Reduce dimensions once more.
    if len(data) > TARGET_BYTES:
        img.thumbnail((1280, 1280), Image.Resampling.LANCZOS)
        for quality in (76, 68, 60, 52):
            data = _encode_jpeg(img, quality)
            if len(data) <= TARGET_BYTES:
                break

    return data, img.size


def add_text(slide, text, x, y, w, h, size=18, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_header(slide, outlet_name, assets: Path, slide_no=None):
    slide.shapes.add_picture(str(assets / "pso_logo.png"), Inches(0.35), Inches(0.18), height=Inches(0.52))
    add_text(slide, outlet_name, Inches(7.0), Inches(0.22), Inches(6.0), Inches(0.4),
             size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(13.333), Inches(0.045))
    line.fill.solid(); line.fill.fore_color.rgb = GREEN; line.line.fill.background()
    line2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), Inches(8.9), Inches(0.045))
    line2.fill.solid(); line2.fill.fore_color.rgb = BLUE; line2.line.fill.background()

    if slide_no is not None:
        add_text(slide, str(slide_no), Inches(12.55), Inches(7.08), Inches(0.4), Inches(0.25),
                 size=9, color=RGBColor(100, 100, 100), align=PP_ALIGN.RIGHT)


def add_picture_cover(slide, data, x, y, w, h):
    with Image.open(BytesIO(data)) as opened:
        img = opened.convert("RGB")
        iw, ih = img.size
        target = float(w) / float(h)
        ratio = iw / ih
        if ratio > target:
            nw = int(ih * target)
            left = max((iw - nw) // 2, 0)
            img = img.crop((left, 0, left + nw, ih))
        else:
            nh = int(iw / target)
            top = max((ih - nh) // 2, 0)
            img = img.crop((0, top, iw, top + nh))
        out = BytesIO()
        img.save(out, format="JPEG", quality=88, optimize=True)
    out.seek(0)
    slide.shapes.add_picture(out, x, y, w, h)


def add_desc(slide, text, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06); tf.margin_top = Inches(0.04)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.name = "Aptos"; r.font.size = Pt(12); r.font.color.rgb = DARK
    return box


def build_ppt(outlet, improvements, assets_dir):
    assets = Path(assets_dir)
    for required in ("pso_logo.png", "cover.jpg"):
        if not (assets / required).exists():
            raise FileNotFoundError(f"Missing asset: {required}")
    if not improvements:
        raise ValueError("At least one improvement is required.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Cover
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(assets / "cover.jpg"), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.shapes.add_picture(str(assets / "pso_logo.png"), Inches(0.35), Inches(0.18), height=Inches(0.62))
    add_text(slide, outlet["name"], Inches(7.0), Inches(0.22), Inches(5.95), Inches(0.4),
             size=12, bold=True, color=BLUE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.45), Inches(6.1), Inches(3.2))
    panel.fill.solid(); panel.fill.fore_color.rgb = WHITE; panel.fill.transparency = 8; panel.line.fill.background()
    add_text(slide, outlet["name"], Inches(1.0), Inches(1.95), Inches(5.4), Inches(0.7), size=28, bold=True)
    add_text(slide, outlet["code"], Inches(1.0), Inches(2.72), Inches(5.4), Inches(0.55), size=22, bold=True)
    add_text(slide, "OUTLET IMPROVEMENT REPORT", Inches(1.0), Inches(3.38), Inches(5.4), Inches(0.6), size=18, bold=True, color=BLUE)

    details = []
    if outlet.get("presented_by"):
        details.append("Presented by: " + outlet["presented_by"])
    if outlet.get("designation"):
        details.append("Designation: " + outlet["designation"])
    if details:
        add_text(slide, "\n".join(details), Inches(1.0), Inches(4.05), Inches(5.4), Inches(0.7), size=11)

    # Improvement slides
    for idx, item in enumerate(improvements, 1):
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = WHITE; bg.line.fill.background()
        add_header(slide, outlet["name"], assets, idx)
        add_text(slide, f"{idx}. {item['heading']}", Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.55),
                 size=23, bold=True, color=BLUE)

        left_x = Inches(0.55); right_x = Inches(6.85); top = Inches(1.72)
        img_w = Inches(5.95); img_h = Inches(3.05)
        entries = [
            (left_x, "BEFORE", item["before"], item["before_desc"]),
            (right_x, "AFTER", item["after"], item["after_desc"]),
        ]
        for x, label, data, desc in entries:
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, img_w, Inches(4.95))
            card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = RGBColor(220, 226, 232)
            add_text(slide, label, x + Inches(0.15), top + Inches(0.12), Inches(1.2), Inches(0.35),
                     size=12, bold=True, color=GREEN)
            add_picture_cover(slide, data, x + Inches(0.15), top + Inches(0.52), img_w - Inches(0.3), img_h)
            add_desc(slide, desc, x + Inches(0.15), top + Inches(3.72), img_w - Inches(0.3), Inches(1.0))

        add_text(slide, "Pakistan State Oil Company Limited", Inches(0.55), Inches(7.02), Inches(6.0), Inches(0.25),
                 size=8, bold=True, color=RGBColor(90, 90, 90))

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
